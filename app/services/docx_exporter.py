import os
import logging
from typing import Optional
from app.services.reconstruction import DocumentDOM, BlockType

logger = logging.getLogger(__name__)


class DOCXExporter:
    """
    Exports DocumentDOM into Microsoft Word (.docx) format.
    Maps DOM Headings, Paragraphs, and Tables to Word elements.
    """

    @staticmethod
    def export_to_docx(dom: DocumentDOM, output_docx_path: str) -> str:
        os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)), exist_ok=True)
        
        try:
            import docx
            doc = docx.Document()
            doc.add_heading(dom.title, level=0)

            for page in dom.pages:
                for block in page.blocks:
                    content = block.content.strip()
                    if not content and not block.table_html:
                        continue

                    if block.block_type == BlockType.TITLE:
                        doc.add_heading(content, level=1)
                    elif block.block_type == BlockType.HEADING:
                        doc.add_heading(content, level=2)
                    elif block.block_type == BlockType.TABLE:
                        p = doc.add_paragraph()
                        p.add_run(f"[Table Content]\n{content}").bold = True
                    else:
                        doc.add_paragraph(content)

            doc.save(output_docx_path)
            logger.info(f"Successfully exported DOCX using python-docx: {output_docx_path}")
            return output_docx_path

        except ImportError:
            # Fallback text/docx placeholder if python-docx not installed
            fallback_txt = output_docx_path.replace(".docx", ".txt")
            with open(fallback_txt, "w", encoding="utf-8") as f:
                f.write(f"# {dom.title}\n\n")
                for page in dom.pages:
                    for block in page.blocks:
                        f.write(f"{block.content}\n\n")
            logger.info(f"python-docx not found, saved formatted document to: {fallback_txt}")
            return fallback_txt
