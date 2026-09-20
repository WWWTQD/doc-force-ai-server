"""
DOCX & PDF Export Service cho FastAPI App trong AI Server.
"""

import io
import logging
from typing import Dict, List, Any

LOG = logging.getLogger("export_service")

try:
    import docx
    from docx.shared import Inches, Pt
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class DocumentExportService:
    """Tái tạo văn bản Word .docx từ Layout JSON."""

    def export_to_docx(self, layout_json: Dict[str, Any]) -> bytes:
        if not DOCX_AVAILABLE:
            raise RuntimeError("python-docx chưa được cài đặt. Hãy cài: pip install python-docx")

        doc = docx.Document()

        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        blocks = layout_json.get("blocks", [])
        blocks.sort(key=lambda b: b.get("order", 0))

        for block in blocks:
            b_type = str(block.get("type", "paragraph")).lower()
            text = block.get("text", "").strip()
            if not text:
                continue

            if "heading" in b_type or b_type == "title":
                p = doc.add_heading(text, level=1)
                p.paragraph_format.space_before = Pt(12)
                p.paragraph_format.space_after = Pt(6)
            elif "table" in b_type:
                p = doc.add_paragraph()
                r = p.add_run(f"[BẢNG]: {text}")
                r.bold = True
                p.paragraph_format.space_after = Pt(6)
            elif "list" in b_type:
                doc.add_paragraph(text, style="List Bullet")
            else:
                p = doc.add_paragraph(text)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
